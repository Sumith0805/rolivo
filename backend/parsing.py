import io
import re

from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader

MAX_BYTES = 2 * 1024 * 1024
MAX_CHARS = 20000


def clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 80:
        raise HTTPException(
            status_code=400,
            detail="There is very little text here. If your PDF is a scan or an image, paste the text instead.",
        )
    return text[:MAX_CHARS]


def text_from_upload(filename: str, content: bytes) -> str:
    name = (filename or "").lower()
    if len(content) > MAX_BYTES:
        raise HTTPException(status_code=400, detail="That file is too large. Keep it under 2 MB.")
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        elif name.endswith(".docx"):
            doc = Document(io.BytesIO(content))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(parts)
        elif name.endswith(".txt"):
            text = content.decode("utf-8", errors="ignore")
        else:
            raise HTTPException(status_code=400, detail="Upload a PDF, a Word (.docx) or a text file.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read that file. Try another one or paste the text.")
    return clean(text)


def find_contacts(text: str) -> dict:
    email = re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", text)
    phone = re.search(r"\+?\d[\d\s().-]{8,}\d", text)
    links = re.findall(r"(?:https?://|www\.)\S+|(?:linkedin\.com|github\.com)/\S+", text)
    return {
        "email": email.group(0) if email else "",
        "phone": phone.group(0).strip() if phone else "",
        "links": sorted({l.rstrip(".,;)") for l in links})[:6],
    }