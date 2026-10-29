"""Reads the structure and look of an uploaded resume, so the rewrite can come back in the same layout.

This is an approximation: it copies the section order, the heading wording, the font family,
the header alignment and the accent color. It does not copy the exact design.
"""
import io
import re
from collections import Counter

SYNONYMS = {
    "summary": {"summary", "professional summary", "profile", "professional profile", "objective", "career objective",
                "about me", "about", "personal statement", "executive summary"},
    "skills": {"skills", "technical skills", "key skills", "core skills", "core competencies", "skills and tools",
               "tools and technologies", "technologies", "skills and abilities", "areas of expertise", "competencies"},
    "experience": {"experience", "work experience", "professional experience", "employment", "employment history",
                   "work history", "internships", "internship experience", "internship", "career history"},
    "projects": {"projects", "academic projects", "personal projects", "key projects", "selected projects", "project work"},
    "education": {"education", "academic background", "academics", "qualifications", "educational qualifications",
                  "academic qualifications", "education and training"},
    "certifications": {"certifications", "certificates", "licenses", "licenses and certifications", "courses",
                       "training", "certifications and courses", "professional development"},
    "achievements": {"achievements", "awards", "honors", "honours", "accomplishments", "awards and achievements",
                     "extracurricular", "extracurricular activities", "activities", "publications", "leadership",
                     "positions of responsibility", "volunteer experience", "volunteering"},
}
_LOOKUP = {name: key for key, names in SYNONYMS.items() for name in names}
SERIF_HINTS = ("times", "georgia", "garamond", "cambria", "palatino", "book antiqua", "century", "constantia",
               "baskerville", "minion", "serif")


def _norm(line):
    t = re.sub(r"[^a-z& ]", " ", line.lower()).replace("&", " and ")
    return re.sub(r"\s+", " ", t).strip()


def find_headings(text):
    """[(key, original line)] in the order the sections appear."""
    found, seen = [], set()
    for raw in str(text).splitlines():
        line = raw.strip().strip(":").strip()
        if not line or len(line) > 45:
            continue
        key = _LOOKUP.get(_norm(line))
        if key and key not in seen:
            seen.add(key)
            found.append((key, line))
    return found


def _tidy_label(line, caps):
    label = re.sub(r"\s+", " ", line).strip()
    if caps or label.isupper():
        label = label.title().replace(" And ", " and ")
    return label[:30]


def _docx_style(data, headings):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document(io.BytesIO(data))
    fonts, sizes = Counter(), Counter()
    default_font = doc.styles["Normal"].font.name
    for p in doc.paragraphs:
        for r in p.runs:
            n = len(r.text.strip())
            if not n:
                continue
            fonts[(r.font.name or p.style.font.name or default_font or "")] += n
            if r.font.size:
                sizes[round(r.font.size.pt * 2) / 2] += n
    style = {}
    font = (fonts.most_common(1) or [("", 0)])[0][0] or default_font or ""
    style["font"] = "serif" if any(h in font.lower() for h in SERIF_HINTS) else "sans"
    if sizes:
        body = sizes.most_common(1)[0][0]
        style["size"] = min(max(body, 9), 12)
    first = next((p for p in doc.paragraphs if p.text.strip()), None)
    style["header"] = "center" if first is not None and first.alignment == WD_ALIGN_PARAGRAPH.CENTER else "left"

    heading_texts = {line.strip().strip(":").strip() for _, line in headings}
    color, ruled, shaded = None, 0, 0
    for p in doc.paragraphs:
        if p.text.strip().strip(":").strip() not in heading_texts:
            continue
        ppr = p._p.pPr
        if ppr is not None and ppr.find(qn("w:pBdr")) is not None:
            ruled += 1
        if ppr is not None and ppr.find(qn("w:shd")) is not None:
            shaded += 1
        for r in p.runs:
            try:
                rgb = r.font.color.rgb if r.font.color is not None and r.font.color.type is not None else None
            except Exception:
                rgb = None
            if rgb is not None and color is None:
                color = "#" + str(rgb).lower()
    if color:
        r_, g_, b_ = (int(color[i:i + 2], 16) for i in (1, 3, 5))
        if max(r_, g_, b_) < 40 or min(r_, g_, b_) > 235:
            color = None
    n = max(len(headings), 1)
    if shaded >= n / 2:
        style["heading"] = "bar"
    elif ruled >= n / 2:
        style["heading"] = "rule"
    elif color:
        style["heading"] = "accent"
    else:
        style["heading"] = "plain"
    if color:
        style["accent"] = color
    return style


def detect_layout(text, docx_bytes=None):
    """A spec for the person's own layout, or None when the sections cannot be read."""
    headings = find_headings(text)
    if len(headings) < 2:
        return None
    caps = sum(1 for _, line in headings if line.isupper()) >= len(headings) / 2
    spec = {
        "id": "yours",
        "order": [k for k, _ in headings],
        "headings": {k: _tidy_label(line, caps) for k, line in headings},
        "caps": caps,
        "heading": "rule",
        "header": "left",
    }
    if docx_bytes:
        try:
            spec.update(_docx_style(docx_bytes, headings))
        except Exception:
            pass  # a damaged or unusual file still gives us the order and wording
    return spec
