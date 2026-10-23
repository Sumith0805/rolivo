import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from auth import current_user
from db import User
from role_profiles import TEMPLATES

router = APIRouter(prefix="/export", tags=["export"])

PLACEHOLDER = "[add a number]"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
AFTER_RULE = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
    "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
    "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


class ExportBody(BaseModel):
    resume: dict
    template: str = "classic"
    allow_placeholders: bool = False


def _s(x):
    return x.strip() if isinstance(x, str) else ""


def _lst(x):
    return x if isinstance(x, list) else []


def _rgb(hex_color):
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_font(font, name):
    font.name = name
    rpr = font._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _run(par, text, size, font, bold=False, italic=False, color=None):
    run = par.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    _set_font(run.font, font)
    if color is not None:
        run.font.color.rgb = color
    return run


def _rule(par, color_hex):
    ppr = par._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color_hex.lstrip("#"))
    borders.append(bottom)
    ppr.insert_element_before(borders, *AFTER_RULE)


def count_placeholders(res):
    texts = [_s(res.get("summary"))]
    for kind in ("experience", "projects"):
        for e in _lst(res.get(kind)):
            if isinstance(e, dict):
                texts += [_s(b) for b in _lst(e.get("bullets"))]
    return sum(t.count(PLACEHOLDER) for t in texts)


def build_docx(res, tpl):
    tight = tpl["density"] == "tight"
    font, accent_hex = tpl["font"], tpl["accent"]
    accent = _rgb(accent_hex)
    ink = RGBColor(0x1A, 0x1A, 0x1A)
    body = 10 if tight else 10.5
    gap = 5 if tight else 9

    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.title = (_s(res.get("name")) + " Resume").strip()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    margin = Inches(0.6 if tight else 0.75)
    sec.left_margin = sec.right_margin = margin
    sec.top_margin = sec.bottom_margin = Inches(0.55 if tight else 0.7)
    width = sec.page_width - 2 * margin

    normal = doc.styles["Normal"]
    _set_font(normal.font, font)
    normal.font.size = Pt(body)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0 if tight else 1.08
    bullet = doc.styles["List Bullet"]
    bullet.paragraph_format.space_after = Pt(1 if tight else 2)

    name = _s(res.get("name"))
    if name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _run(p, name, 20 if not tight else 18, font, bold=True, color=accent)
    if _s(res.get("headline")):
        _run(doc.add_paragraph(), _s(res.get("headline")), body + 0.5, font, color=ink)
    c = res.get("contact") if isinstance(res.get("contact"), dict) else {}
    bits = [_s(c.get("email")), _s(c.get("phone")), _s(c.get("location"))]
    bits += [_s(x) for x in _lst(c.get("links"))]
    line = "  |  ".join(b for b in bits if b)
    if line:
        _run(doc.add_paragraph(), line, body - 0.5, font, color=ink)

    def heading(text):
        p = doc.add_paragraph()
        _rule(p, accent_hex)
        p.paragraph_format.space_before = Pt(gap)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        _run(p, text.upper(), body + 0.5, font, bold=True, color=accent)

    def title_line(left_bold, left_rest, right):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        _run(p, left_bold, body, font, bold=True, color=ink)
        if left_rest:
            _run(p, left_rest, body, font, color=ink)
        if right:
            _run(p, "\t" + right, body, font, color=ink)

    def bullets(items):
        for b in items:
            if _s(b):
                p = doc.add_paragraph(style="List Bullet")
                _run(p, _s(b), body, font, color=ink)

    def section(key):
        if key == "summary" and _s(res.get("summary")):
            heading("Summary")
            _run(doc.add_paragraph(), _s(res.get("summary")), body, font, color=ink)
        elif key == "skills" and _lst(res.get("skills")):
            heading("Skills")
            for g in _lst(res.get("skills")):
                if isinstance(g, dict) and _lst(g.get("items")):
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(1)
                    _run(p, _s(g.get("group")) + ": ", body, font, bold=True, color=ink)
                    items = ", ".join(_s(i) for i in _lst(g.get("items")) if _s(i))
                    _run(p, items, body, font, color=ink)
        elif key == "experience" and _lst(res.get("experience")):
            heading("Experience")
            for e in _lst(res.get("experience")):
                if isinstance(e, dict):
                    title = _s(e.get("title"))
                    parts = [_s(e.get("company")), _s(e.get("location"))]
                    where = ", ".join(x for x in parts if x)
                    dates = _s(e.get("dates"))
                    if title:
                        title_line(title, (", " + where) if where else "", dates)
                    else:
                        title_line(where, "", dates)
                    bullets(_lst(e.get("bullets")))
        elif key == "projects" and _lst(res.get("projects")):
            heading("Projects")
            for e in _lst(res.get("projects")):
                if isinstance(e, dict):
                    tech = " | " + _s(e.get("tech")) if _s(e.get("tech")) else ""
                    title_line(_s(e.get("name")), tech, "")
                    bullets(_lst(e.get("bullets")))
        elif key == "education" and _lst(res.get("education")):
            heading("Education")
            for e in _lst(res.get("education")):
                if isinstance(e, dict):
                    school = ", " + _s(e.get("school")) if _s(e.get("school")) else ""
                    title_line(_s(e.get("degree")), school, _s(e.get("dates")))
                    if _s(e.get("details")):
                        _run(doc.add_paragraph(), _s(e.get("details")), body, font, color=ink)
        elif key == "certifications" and _lst(res.get("certifications")):
            heading("Certifications")
            bullets([str(x) for x in _lst(res.get("certifications"))])

    for key in tpl["order"]:
        section(key)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


@router.post("/docx")
def export_docx(body: ExportBody, user: User = Depends(current_user)):
    tpl = TEMPLATES.get(body.template) or TEMPLATES["classic"]
    n = count_placeholders(body.resume)
    if n and not body.allow_placeholders:
        raise HTTPException(
            status_code=400,
            detail=f"Replace or delete the {n} {PLACEHOLDER} markers before exporting.",
        )
    data = build_docx(body.resume, tpl)
    safe = re.sub(r"[^A-Za-z0-9]+", "_", _s(body.resume.get("name")) or "Resume").strip("_")
    headers = {"Content-Disposition": f'attachment; filename="{safe}_Resume.docx"'}
    return StreamingResponse(io.BytesIO(data), media_type=DOCX_TYPE, headers=headers)